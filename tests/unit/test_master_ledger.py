"""
model.master_ledger（UI 非依存）のユニットテスト。

streamlit に依存しないため app.py をインポートせず、コアを直接検証する
（tests/unit/test_pairing.py と同じ方針）。

実行:
    cd DXF-diff-manager
    python -m tests.unit.test_master_ledger
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

import pandas as pd

from model.master_ledger import (
    load_parent_child_master,
    update_parent_child_master,
    create_empty_master_df,
    save_master_to_bytes,
    make_dataframe_arrow_compatible,
    parse_master_filename,
    create_empty_drawing_list_df,
    load_drawing_list,
    update_drawing_list,
    DRAWING_LIST_SHEET_NAME,
    MASTER_SHEET_NAME,
)


# --- create_empty_master_df ---

def test_create_empty_master_df_has_required_columns():
    df = create_empty_master_df()
    assert list(df.columns) == [
        'Child', 'Parent', 'Relation', 'Title', 'Subtitle', 'Recorded Date', 'Note',
        'Deleted Entities', 'Added Entities', 'Diff Entities', 'Unchanged Entities', 'Total Entities',
    ]
    assert len(df) == 0


# --- update_parent_child_master ---

def test_update_parent_child_master_adds_new_record():
    master_df = create_empty_master_df()
    pair = {
        'main_drawing': 'B1', 'source_drawing': 'A1', 'relation': '流用',
        'title': 'T', 'subtitle': 'S',
        'entity_counts': {
            'deleted_entities': 1, 'added_entities': 2,
            'diff_entities': 3, 'unchanged_entities': 4, 'total_entities': 5,
        },
    }
    updated, added_count = update_parent_child_master(master_df, [pair])
    assert added_count == 1
    row = updated[updated['Child'] == 'B1'].iloc[0]
    assert row['Parent'] == 'A1'
    assert row['Deleted Entities'] == 1
    assert row['Total Entities'] == 5


def test_update_parent_child_master_skips_pair_without_child():
    master_df = create_empty_master_df()
    pair = {'main_drawing': None, 'source_drawing': 'A1'}
    updated, added_count = update_parent_child_master(master_df, [pair])
    assert added_count == 0
    assert len(updated) == 0


def test_update_parent_child_master_no_futurewarning_setting_na_on_numeric_column():
    """完全新規図面の先行登録（entity_counts 未確定 → 数値カラムが NaN のみで
    float64 として残る）の後、create_diff_zip() 側の2回目の update 呼び出しで
    実際のエンティティ数を書き込む際、同じ行の 'Deleted/Diff/Unchanged Entities'
    に "n/a" 文字列を代入する（既存レコード更新パス、model/master_ledger.py の
    元の行141）。台帳を Excel 経由で読み込んだ場合など、この時点でカラムが
    まだ float64 のままだと pandas FutureWarning（将来 TypeError 化予定）が出て
    いた。警告を例外に昇格させ、出ないことを確認する回帰テスト。"""
    import warnings

    master_df = pd.DataFrame({
        'Child': ['B1'], 'Parent': ['none'], 'Relation': ['完全新規図面'],
        'Title': [None], 'Subtitle': [None], 'Recorded Date': [None], 'Note': [None],
        'Deleted Entities': [float('nan')], 'Added Entities': [float('nan')],
        'Diff Entities': [float('nan')], 'Unchanged Entities': [float('nan')],
        'Total Entities': [float('nan')],
    })
    assert master_df['Deleted Entities'].dtype != object  # 前提: NaNのみで float64

    brand_new_pair = {
        'main_drawing': 'B1', 'source_drawing': None, 'relation': '完全新規図面',
        'title': None, 'subtitle': None,
        'entity_counts': {'added_entities': 6, 'total_entities': 6},
    }
    with warnings.catch_warnings():
        warnings.simplefilter('error', FutureWarning)
        updated, added_count = update_parent_child_master(master_df, [brand_new_pair])
    assert added_count == 0  # 既存レコードの更新
    row = updated[updated['Child'] == 'B1'].iloc[0]
    assert row['Deleted Entities'] == 'n/a'
    assert row['Added Entities'] == 6


def test_update_parent_child_master_existing_record_relation_changed_suffix():
    master_df = create_empty_master_df()
    first = {'main_drawing': 'B1', 'source_drawing': 'A1', 'relation': '流用'}
    updated, _ = update_parent_child_master(master_df, [first])
    second = {'main_drawing': 'B1', 'source_drawing': 'A1', 'relation': 'RevUp'}
    updated, added_count = update_parent_child_master(updated, [second])
    assert added_count == 0  # 既存レコードの更新（新規追加ではない）
    row = updated[updated['Child'] == 'B1'].iloc[0]
    assert row['Relation'] == 'RevUp-changed'


# --- load_parent_child_master ---

def test_load_parent_child_master_missing_required_column(tmp_path):
    path = tmp_path / "master.xlsx"
    pd.DataFrame({'Child': ['B1']}).to_excel(path, index=False)  # Parent列なし
    df, error = load_parent_child_master(str(path))
    assert df is None
    assert 'Parent' in error


def test_load_parent_child_master_success(tmp_path):
    path = tmp_path / "master.xlsx"
    pd.DataFrame({'Child': ['B1'], 'Parent': ['A1']}).to_excel(path, index=False)
    df, error = load_parent_child_master(str(path))
    assert error is None
    assert len(df) == 1


def test_load_parent_child_master_finds_data_sheet_when_first_sheet_has_no_child_column(tmp_path):
    """save_master_to_bytes() が出力する台帳（Summaryシートが先頭）を再アップロード
    しても、Child/Parent 列を持つシート（Diff List）を自動で見つけて読み込める。

    実データ（ME24-9001-0_ZM00_405.xlsx）で「必須カラム 'Child' が見つかりません」
    と誤って失敗していた不具合の回帰テスト: 先頭シート（Summary）を無条件に読んで
    いたため、Child/Parent 列を持つ実データシート（Diff List）が無視されていた。
    """
    path = tmp_path / "master.xlsx"
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        # Summary シート（Child/Parent 列を持たない、統計情報のみ）を先頭に作成
        pd.DataFrame({'エンティティ統計': ['削除図形 総数'], 'Unnamed: 1': [10]}).to_excel(
            writer, sheet_name='Summary', index=False)
        # Diff List シート（実データ）を2番目に作成
        pd.DataFrame({'Child': ['B1'], 'Parent': ['A1']}).to_excel(
            writer, sheet_name='Diff List', index=False)

    df, error = load_parent_child_master(str(path))
    assert error is None
    assert df is not None
    assert list(df['Child']) == ['B1']
    assert list(df['Parent']) == ['A1']


def test_load_parent_child_master_finds_data_sheet_named_master(tmp_path):
    """2026-08改名後の現行フォーマット（データシート名 'Master'）も、旧名
    'Diff List' と同様に列の有無で見つけて読み込める（シート名を問わない設計の確認）。"""
    path = tmp_path / "master.xlsx"
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        pd.DataFrame({'エンティティ統計': ['削除図形 総数'], 'Unnamed: 1': [10]}).to_excel(
            writer, sheet_name='Summary', index=False)
        pd.DataFrame({'Child': ['B1'], 'Parent': ['A1']}).to_excel(
            writer, sheet_name=MASTER_SHEET_NAME, index=False)

    df, error = load_parent_child_master(str(path))
    assert error is None
    assert df is not None
    assert list(df['Child']) == ['B1']
    assert list(df['Parent']) == ['A1']


def test_save_master_to_bytes_round_trip_reloads_correctly(tmp_path):
    """save_master_to_bytes() の出力をそのまま load_parent_child_master() で
    再読み込みできる（エクスポート→再アップロードの往復を保証する）。"""
    master_df = create_empty_master_df()
    master_df.loc[0] = {
        'Child': 'B1', 'Parent': 'A1', 'Relation': '流用',
        'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
        'Deleted Entities': 1, 'Added Entities': 2, 'Diff Entities': 3,
        'Unchanged Entities': 4, 'Total Entities': 5,
    }
    data = save_master_to_bytes(master_df, mode='auto')

    path = tmp_path / "roundtrip.xlsx"
    path.write_bytes(data)

    df, error = load_parent_child_master(str(path))
    assert error is None
    assert df is not None
    assert list(df['Child']) == ['B1']
    assert list(df['Parent']) == ['A1']


def test_uploaded_master_merges_correctly_across_all_pairing_modes(tmp_path):
    """Step0でアップロードした台帳（Summary+Diff List形式）が、Step1のどの
    ペアリング方式（Type A/B の RevUp・流用、Type C のペアリスト）で得られた
    ペアとも正しく合流する（新規追加は重複なく、既存行は上書き更新される）。

    dev-workflow スキルの選択肢組み合わせ表で「台帳アップロード × Step1モード」を
    影響あり→要確認と判定した組み合わせの回帰テスト。update_parent_child_master()
    は pairs のスキーマ（pairing.py で全モード共通と規定）のみに依存し mode 分岐を
    持たないため構造的には安全なはずだが、実際にアップロード経由で読み込んだ
    DataFrame に対して確認する。
    """
    # Step0でアップロードされる形式（Summaryシートが先頭）の台帳を用意
    path = tmp_path / "uploaded_master.xlsx"
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        pd.DataFrame({'エンティティ統計': ['削除図形 総数'], 'Unnamed: 1': [0]}).to_excel(
            writer, sheet_name='Summary', index=False)
        pd.DataFrame({'Child': ['EXIST-CHILD'], 'Parent': ['EXIST-PARENT'], 'Relation': ['流用']}).to_excel(
            writer, sheet_name='Diff List', index=False)

    master_df, error = load_parent_child_master(str(path))
    assert error is None

    # Type A/B の RevUp・流用、Type C のペアリスト、それぞれ新規追加のケース
    for relation in ('RevUp', '流用', 'ペアリスト'):
        pair = {
            'main_drawing': f'NEW-{relation}', 'source_drawing': f'OLD-{relation}',
            'relation': relation, 'title': 'T', 'subtitle': 'S',
        }
        master_df, added_count = update_parent_child_master(master_df, [pair])
        assert added_count == 1, f"{relation} ペアが新規追加されなかった"

    assert len(master_df) == 4  # 既存1件 + 新規3件

    # 既存行の更新（Type C のペアリストで同じ Child/Parent が再検出されたケース）
    update_pair = {
        'main_drawing': 'EXIST-CHILD', 'source_drawing': 'EXIST-PARENT',
        'relation': 'ペアリスト', 'title': 'T2', 'subtitle': 'S2',
    }
    master_df, added_count = update_parent_child_master(master_df, [update_pair])
    assert added_count == 0, "既存行が更新ではなく新規追加されてしまった（重複）"
    assert len(master_df) == 4  # 行数は増えない

    match = master_df[(master_df['Child'] == 'EXIST-CHILD') & (master_df['Parent'] == 'EXIST-PARENT')]
    assert len(match) == 1
    # Relation は「新しい値+-changed」で記録される（既存の
    # test_update_parent_child_master_existing_record_relation_changed_suffix と同じ仕様）
    assert match.iloc[0]['Relation'] == 'ペアリスト-changed'


def test_load_parent_child_master_no_matching_sheet_returns_error(tmp_path):
    """どのシートにも Child/Parent 列が無い場合は、従来どおりエラーを返す
    （先頭シートを対象にエラーメッセージを出す後方互換の挙動）。"""
    path = tmp_path / "master.xlsx"
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        pd.DataFrame({'foo': [1]}).to_excel(writer, sheet_name='Sheet1', index=False)
        pd.DataFrame({'bar': [2]}).to_excel(writer, sheet_name='Sheet2', index=False)

    df, error = load_parent_child_master(str(path))
    assert df is None
    assert 'Child' in error


# --- save_master_to_bytes ---

def test_save_master_to_bytes_returns_nonempty_excel():
    master_df = create_empty_master_df()
    data = save_master_to_bytes(master_df, mode='auto')
    assert isinstance(data, bytes) and len(data) > 0


def test_save_master_to_bytes_sorts_diff_list_by_child():
    """Master シートは Child 列の昇順（ABC順）でソートされる。"""
    master_df = create_empty_master_df()
    for i, child in enumerate(['EE3273-608-32B', 'EE3273-608-24B', 'DE5313-008-02A']):
        master_df.loc[i] = {
            'Child': child, 'Parent': 'none', 'Relation': '完全新規図面',
            'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
            'Deleted Entities': 'n/a', 'Added Entities': 1, 'Diff Entities': 'n/a',
            'Unchanged Entities': 'n/a', 'Total Entities': 1,
        }

    data = save_master_to_bytes(master_df, mode='pair_list')
    diff_list_df = pd.read_excel(pd.io.common.BytesIO(data), sheet_name=MASTER_SHEET_NAME)
    assert list(diff_list_df['Child']) == ['DE5313-008-02A', 'EE3273-608-24B', 'EE3273-608-32B']
    # 元の master_df は変更されない（呼び出し元の順序に副作用を与えない）
    assert list(master_df['Child']) == ['EE3273-608-32B', 'EE3273-608-24B', 'DE5313-008-02A']


def test_save_master_to_bytes_summary_has_no_drawing_statistics_section():
    """Summaryシートから「図面統計」欄（見出し＋流用先図面総数〜新規作成率の
    全行）が削除されている（2026-09）。Ledger-merger 等の下流に「流用先図面総数」
    のような分母を渡せなくなる仕様変更で、ユーザーが下流影響を承知のうえで
    全削除を選択した。エンティティ統計（削除/追加/変更/変更なし/総数・図形変更率）
    は引き続き出力される。
    """
    master_df = create_empty_master_df()
    master_df.loc[0] = {
        'Child': 'B1', 'Parent': 'A1', 'Relation': '流用',
        'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
        'Deleted Entities': 1, 'Added Entities': 2, 'Diff Entities': 3,
        'Unchanged Entities': 4, 'Total Entities': 5,
    }
    master_df.loc[1] = {
        'Child': 'B2', 'Parent': 'none', 'Relation': '完全新規図面',
        'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
        'Deleted Entities': 'n/a', 'Added Entities': 9, 'Diff Entities': 'n/a',
        'Unchanged Entities': 'n/a', 'Total Entities': 9,
    }
    data = save_master_to_bytes(master_df, mode='auto')
    xl = pd.ExcelFile(pd.io.common.BytesIO(data))
    summary_df = pd.read_excel(xl, sheet_name='Summary', header=None)
    labels = summary_df[0].dropna().tolist()

    removed_labels = (
        '図面統計', '流用先図面総数', 'アップロード図面総数',
        '差分抽出ペア数', '完全新規図面数', '流用率 [%]', '新規作成率 [%]',
    )
    for label in removed_labels:
        assert label not in labels, f"「図面統計」削除後も残っているラベル: {label}"

    remaining_labels = (
        'エンティティ統計', '削除図形 総数', '追加図形 総数',
        '変更（追加+削除）図形 総数', '変更なし図形 総数', '図形変更率 [%]',
    )
    for label in remaining_labels:
        assert label in labels, f"エンティティ統計のラベルが失われている: {label}"


def test_save_master_to_bytes_freezes_header_row_for_diff_list_and_drawing_list():
    """Master・Drawing Listのタイトル行（1行目）が固定される。"""
    import openpyxl
    master_df = create_empty_master_df()
    drawing_list_df = create_empty_drawing_list_df()
    data = save_master_to_bytes(master_df, mode='auto', drawing_list_df=drawing_list_df)
    wb = openpyxl.load_workbook(pd.io.common.BytesIO(data))
    assert wb[MASTER_SHEET_NAME].freeze_panes == 'A2'
    assert wb[DRAWING_LIST_SHEET_NAME].freeze_panes == 'A2'


def test_save_master_to_bytes_centers_and_formats_entity_columns():
    """「* Entities」列（整数と'n/a'文字列が混在）は、数値・'n/a'セルとも中央揃え・
    桁区切り書式になる（2026-08 ユーザー指摘: 数値は右揃え、'n/a'は左揃えで
    アラインが揃わず見づらかった）。"""
    import openpyxl
    master_df = create_empty_master_df()
    master_df.loc[0] = {
        'Child': 'B1', 'Parent': 'A1', 'Relation': '流用',
        'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
        'Deleted Entities': 1234, 'Added Entities': 5678, 'Diff Entities': 6912,
        'Unchanged Entities': 100, 'Total Entities': 7012,
    }
    master_df.loc[1] = {
        'Child': 'B2', 'Parent': 'none', 'Relation': '完全新規図面',
        'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
        'Deleted Entities': 'n/a', 'Added Entities': 9999, 'Diff Entities': 'n/a',
        'Unchanged Entities': 'n/a', 'Total Entities': 9999,
    }
    data = save_master_to_bytes(master_df, mode='auto')
    wb = openpyxl.load_workbook(pd.io.common.BytesIO(data))
    ws = wb[MASTER_SHEET_NAME]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    entity_cols = {'Deleted Entities', 'Added Entities', 'Diff Entities',
                   'Unchanged Entities', 'Total Entities'}
    for row in ws.iter_rows(min_row=2, max_row=3):
        for cell in row:
            col_name = header[cell.column - 1]
            if col_name in entity_cols:
                assert cell.alignment.horizontal == 'center', \
                    f"{col_name}={cell.value!r} が中央揃えでない: {cell.alignment.horizontal}"
                assert cell.number_format == '#,##0', \
                    f"{col_name}={cell.value!r} の桁区切り書式が適用されていない: {cell.number_format}"
            else:
                assert cell.alignment.horizontal is None, \
                    f"対象外の列{col_name}にまで中央揃えが適用されている"


def test_save_master_to_bytes_handles_na_entity_strings():
    """完全新規図面の 'n/a' 文字列が混在してもサマリー合計でエラーにならない。"""
    master_df = create_empty_master_df()
    master_df.loc[0] = {
        'Child': 'B1', 'Parent': 'none', 'Relation': '完全新規図面',
        'Title': None, 'Subtitle': None, 'Recorded Date': None, 'Note': None,
        'Deleted Entities': 'n/a', 'Added Entities': 10, 'Diff Entities': 'n/a',
        'Unchanged Entities': 'n/a', 'Total Entities': 10,
    }
    data = save_master_to_bytes(master_df, mode='pair_list')
    assert isinstance(data, bytes) and len(data) > 0


# --- make_dataframe_arrow_compatible ---

def test_make_dataframe_arrow_compatible_mixed_entity_columns():
    """'n/a' と整数が混在するエントリ数カラムが pyarrow でシリアライズ可能になる。

    完全新規図面の行（'n/a'）と通常ペアの行（整数）が混在した台帳をそのまま
    st.dataframe に渡すと pyarrow が変換に失敗して警告を出す。表示用コピーで
    混在カラムを文字列統一することで Arrow 互換になることを検証する。
    """
    import datetime
    pa = __import__('pyarrow')

    master_df = create_empty_master_df()
    master_df.loc[0] = {
        'Child': 'B1', 'Parent': 'A1', 'Relation': '流用',
        'Title': 'T', 'Subtitle': 'S',
        'Recorded Date': datetime.datetime(2026, 7, 8), 'Note': None,
        'Deleted Entities': 10, 'Added Entities': 20, 'Diff Entities': 30,
        'Unchanged Entities': 40, 'Total Entities': 100,
    }
    master_df.loc[1] = {
        'Child': 'B2', 'Parent': 'none', 'Relation': '完全新規図面',
        'Title': 'T2', 'Subtitle': None,
        'Recorded Date': datetime.datetime(2026, 7, 8), 'Note': None,
        'Deleted Entities': 'n/a', 'Added Entities': 5, 'Diff Entities': 'n/a',
        'Unchanged Entities': 'n/a', 'Total Entities': 5,
    }

    # 修正前は失敗することを確認（回帰の前提）
    try:
        pa.Table.from_pandas(master_df)
        raise AssertionError("前提が崩れている: 元の混在DataFrameがArrow変換に成功してしまった")
    except pa.lib.ArrowInvalid:
        pass

    display_df = make_dataframe_arrow_compatible(master_df)

    # 修正後は成功する
    pa.Table.from_pandas(display_df)  # 例外が出ないこと

    # 元のDataFrameは変更されない
    assert master_df.loc[0, 'Deleted Entities'] == 10

    # 混在カラムは文字列統一される
    assert display_df.loc[0, 'Deleted Entities'] == '10'
    assert display_df.loc[1, 'Deleted Entities'] == 'n/a'

    # 純粋な整数カラム（全行が数値）は数値のまま維持される
    assert display_df.loc[0, 'Total Entities'] == 100


def test_make_dataframe_arrow_compatible_leaves_clean_columns_untouched():
    """数値のみ・文字列のみの純粋なカラムは変換されない。"""
    df = pd.DataFrame({
        'counts': [1, 2, 3],           # 数値のみ
        'labels': ['a', 'b', 'c'],     # 文字列のみ
    })
    display_df = make_dataframe_arrow_compatible(df)
    assert display_df['counts'].tolist() == [1, 2, 3]  # 数値のまま
    assert display_df['labels'].tolist() == ['a', 'b', 'c']


# --- parse_master_filename ---

def test_parse_master_filename_matches_naming_convention():
    shiban, module, side = parse_master_filename('AA11-1111-1_ZM00_405.xlsx')
    assert (shiban, module, side) == ('AA11-1111-1', 'ZM00', '405')


def test_parse_master_filename_handles_na_module_and_side():
    shiban, module, side = parse_master_filename('AA11-1111-1_na_na.xlsx')
    assert (shiban, module, side) == ('AA11-1111-1', 'na', 'na')


def test_parse_master_filename_returns_none_for_non_matching_name():
    """自由な名前でアップロードされた台帳（命名規則不一致）は (None, None, None)。"""
    assert parse_master_filename('my_master_ledger.xlsx') == (None, None, None)
    assert parse_master_filename(None) == (None, None, None)


def test_parse_master_filename_allows_underscore_suffix():
    """指番_モジュール_サイドの後ろに "_" 区切りで任意の文字列が続いても認識する
    （2026-08、既存台帳の再アップロード時に元のダウンロードファイル名〈例:
    Ledger-merger 側が生成する "..._all.xlsx"〉に接尾辞が付くケースへの対応）。"""
    shiban, module, side = parse_master_filename('AA11-1111-1_ZM00_405_all.xlsx')
    assert (shiban, module, side) == ('AA11-1111-1', 'ZM00', '405')


def test_parse_master_filename_allows_hyphen_suffix():
    """区切り文字は "-" も許容する。"""
    shiban, module, side = parse_master_filename('AA11-1111-1_ZM00_405-all.xlsx')
    assert (shiban, module, side) == ('AA11-1111-1', 'ZM00', '405')


def test_parse_master_filename_rejects_suffix_without_separator():
    """区切り文字なしで直接くっつく形式は非対応のまま
    （サイド直後の文字がサイドの一部か付加文字列かを区別できないため）。"""
    assert parse_master_filename('AA11-1111-1_ZM00_4050.xlsx') == (None, None, None)


# --- create_empty_drawing_list_df ---

def test_create_empty_drawing_list_df_has_required_columns():
    df = create_empty_drawing_list_df()
    assert list(df.columns) == [
        'Sashiban', 'Module', 'Side', 'Child Drawing Number',
        'Parent Drawing Number', 'Title', 'Subtitle', 'Recorded Date',
    ]
    assert len(df) == 0


# --- load_drawing_list ---

def test_load_drawing_list_returns_empty_df_when_sheet_missing(tmp_path):
    """本機能導入前に作成された台帳（Drawing List シートが無い）はエラーにせず空DFを返す。"""
    path = tmp_path / "old_master.xlsx"
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        pd.DataFrame({'Child': ['C1'], 'Parent': ['P1']}).to_excel(
            writer, sheet_name='Diff List', index=False)

    df, error = load_drawing_list(str(path))
    assert error is None
    assert list(df.columns) == list(create_empty_drawing_list_df().columns)
    assert len(df) == 0


def test_load_drawing_list_reads_existing_sheet(tmp_path):
    path = tmp_path / "master_with_drawing_list.xlsx"
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        pd.DataFrame({'Child': ['C1'], 'Parent': ['P1']}).to_excel(
            writer, sheet_name='Diff List', index=False)
        pd.DataFrame({
            'Sashiban': ['AA11-1111-1'], 'Module': ['ZM00'], 'Side': ['405'],
            'Child Drawing Number': ['C1'], 'Parent Drawing Number': ['P1'],
            'Title': ['T'], 'Subtitle': ['S'], 'Recorded Date': ['2026-07-01'],
        }).to_excel(writer, sheet_name=DRAWING_LIST_SHEET_NAME, index=False)

    df, error = load_drawing_list(str(path))
    assert error is None
    assert list(df['Child Drawing Number']) == ['C1']
    assert list(df['Sashiban']) == ['AA11-1111-1']


# --- update_drawing_list ---

def test_update_drawing_list_adds_new_child():
    df = create_empty_drawing_list_df()
    entries = [{'main_drawing': 'C1', 'source_drawing': 'P1', 'title': 'T', 'subtitle': 'S'}]
    updated, added_count = update_drawing_list(df, entries, 'AA11-1111-1', 'ZM00', '405')
    assert added_count == 1
    row = updated[updated['Child Drawing Number'] == 'C1'].iloc[0]
    assert row['Parent Drawing Number'] == 'P1'
    assert row['Sashiban'] == 'AA11-1111-1'
    assert row['Module'] == 'ZM00'
    assert row['Side'] == '405'
    assert row['Title'] == 'T'
    assert row['Subtitle'] == 'S'


def test_update_drawing_list_no_parent_recorded_as_none():
    df = create_empty_drawing_list_df()
    entries = [{'main_drawing': 'C1', 'source_drawing': None, 'title': None, 'subtitle': None}]
    updated, _ = update_drawing_list(df, entries, 'AA11-1111-1', 'ZM00', '405')
    assert updated.iloc[0]['Parent Drawing Number'] == 'none'


def test_update_drawing_list_does_not_overwrite_existing_child():
    """既存のChild Drawing Numberは、Parent/Titleが変わっても上書きしない
    （仕様: 「新規のChild Drawing Numberがあれば追加する」——既存行は保持）。"""
    df = create_empty_drawing_list_df()
    df.loc[0] = {
        'Sashiban': 'AA11-1111-1', 'Module': 'ZM00', 'Side': '405',
        'Child Drawing Number': 'C1', 'Parent Drawing Number': 'OLD-PARENT',
        'Title': 'OLD-TITLE', 'Subtitle': 'OLD-SUB', 'Recorded Date': '2026-01-01',
    }
    entries = [{'main_drawing': 'C1', 'source_drawing': 'NEW-PARENT', 'title': 'NEW-TITLE', 'subtitle': 'NEW-SUB'}]
    updated, added_count = update_drawing_list(df, entries, 'AA11-1111-1', 'ZM00', '405')
    assert added_count == 0
    assert len(updated) == 1
    assert updated.iloc[0]['Parent Drawing Number'] == 'OLD-PARENT'
    assert updated.iloc[0]['Title'] == 'OLD-TITLE'


def test_update_drawing_list_dedups_duplicate_child_within_same_batch():
    """同一バッチ内で同じChildが複数回登場する場合（RevUp+流用の二重登場等）は
    1回のみ追加する（先勝ち）。"""
    df = create_empty_drawing_list_df()
    entries = [
        {'main_drawing': 'C1', 'source_drawing': 'REVUP-PARENT', 'title': 'T1', 'subtitle': None},
        {'main_drawing': 'C1', 'source_drawing': 'DEP-PARENT', 'title': 'T2', 'subtitle': None},
    ]
    updated, added_count = update_drawing_list(df, entries, 'AA11-1111-1', 'ZM00', '405')
    assert added_count == 1
    assert len(updated) == 1
    assert updated.iloc[0]['Parent Drawing Number'] == 'REVUP-PARENT'


def test_update_drawing_list_skips_entries_without_child():
    df = create_empty_drawing_list_df()
    entries = [{'main_drawing': None, 'source_drawing': 'P1', 'title': None, 'subtitle': None}]
    updated, added_count = update_drawing_list(df, entries, 'AA11-1111-1', 'ZM00', '405')
    assert added_count == 0
    assert len(updated) == 0


def test_update_drawing_list_blank_shiban_module_side_when_unresolvable():
    """台帳ファイル名が命名規則に一致しない場合、Sashiban/Module/Sideは空欄で記録する。"""
    df = create_empty_drawing_list_df()
    entries = [{'main_drawing': 'C1', 'source_drawing': 'P1', 'title': None, 'subtitle': None}]
    updated, _ = update_drawing_list(df, entries, None, None, None)
    row = updated.iloc[0]
    assert row['Sashiban'] == ''
    assert row['Module'] == ''
    assert row['Side'] == ''


# --- save_master_to_bytes（Drawing List シート） ---

def test_save_master_to_bytes_writes_drawing_list_sheet_after_diff_list():
    master_df = create_empty_master_df()
    drawing_list_df = create_empty_drawing_list_df()
    drawing_list_df.loc[0] = {
        'Sashiban': 'AA11-1111-1', 'Module': 'ZM00', 'Side': '405',
        'Child Drawing Number': 'C1', 'Parent Drawing Number': 'P1',
        'Title': 'T', 'Subtitle': 'S', 'Recorded Date': None,
    }
    data = save_master_to_bytes(master_df, mode='auto', drawing_list_df=drawing_list_df)
    xl = pd.ExcelFile(pd.io.common.BytesIO(data))
    assert xl.sheet_names == ['Summary', MASTER_SHEET_NAME, DRAWING_LIST_SHEET_NAME]
    dl = pd.read_excel(xl, sheet_name=DRAWING_LIST_SHEET_NAME)
    assert list(dl['Child Drawing Number']) == ['C1']


def test_save_master_to_bytes_drawing_list_defaults_to_empty_when_omitted():
    """drawing_list_df を渡さない既存呼び出し元（後方互換）でも空のDrawing Listシートが出力される。"""
    master_df = create_empty_master_df()
    data = save_master_to_bytes(master_df, mode='auto')
    xl = pd.ExcelFile(pd.io.common.BytesIO(data))
    assert DRAWING_LIST_SHEET_NAME in xl.sheet_names
    dl = pd.read_excel(xl, sheet_name=DRAWING_LIST_SHEET_NAME)
    assert len(dl) == 0
    assert list(dl.columns) == list(create_empty_drawing_list_df().columns)


def test_save_master_to_bytes_sorts_drawing_list_by_child_drawing_number():
    master_df = create_empty_master_df()
    drawing_list_df = create_empty_drawing_list_df()
    for i, child in enumerate(['EE3273-608-32B', 'EE3273-608-24B', 'DE5313-008-02A']):
        drawing_list_df.loc[i] = {
            'Sashiban': 'AA11-1111-1', 'Module': 'ZM00', 'Side': '405',
            'Child Drawing Number': child, 'Parent Drawing Number': 'none',
            'Title': None, 'Subtitle': None, 'Recorded Date': None,
        }
    data = save_master_to_bytes(master_df, mode='auto', drawing_list_df=drawing_list_df)
    dl = pd.read_excel(pd.io.common.BytesIO(data), sheet_name=DRAWING_LIST_SHEET_NAME)
    assert list(dl['Child Drawing Number']) == ['DE5313-008-02A', 'EE3273-608-24B', 'EE3273-608-32B']


def _run_all():
    tests = [v for k, v in sorted(globals().items()) if k.startswith('test_') and callable(v)]
    failures = []
    import tempfile
    for t in tests:
        try:
            if 'tmp_path' in t.__code__.co_varnames[:t.__code__.co_argcount]:
                with tempfile.TemporaryDirectory() as d:
                    from pathlib import Path
                    t(Path(d))
            else:
                t()
            print(f"PASS: {t.__name__}")
        except AssertionError as e:
            failures.append(t.__name__); print(f"FAIL: {t.__name__}\n      {e}")
    print(f"\n{len(tests) - len(failures)}/{len(tests)} passed")
    return 1 if failures else 0


if __name__ == '__main__':
    sys.exit(_run_all())
